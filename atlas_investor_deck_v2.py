"""Generate the Atlas Investor Deck v2 (Polished Edition).

This script builds a PowerPoint presentation summarizing the Atlas investment
model. It mirrors the layout requested in the prompt and can be re-run to
regenerate the deck.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
import io
import matplotlib.pyplot as plt


def add_chart(slide, fig, left=Inches(1), top=Inches(1.5), width=Inches(8)):
    """Render a Matplotlib figure onto the provided slide."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    slide.shapes.add_picture(buf, left, top, width)
    buf.close()


def build_presentation() -> Presentation:
    """Construct the presentation deck."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    navy = RGBColor(10, 35, 66)
    gold = RGBColor(197, 162, 83)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = navy

    text_box = slide.shapes.add_textbox(Inches(1.5), Inches(3), Inches(10), Inches(1))
    text_frame = text_box.text_frame
    paragraph = text_frame.add_paragraph()
    paragraph.text = "Transforming macro chaos into compounding clarity."
    paragraph.font.size = Pt(36)
    paragraph.font.color.rgb = gold

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Atlas Investment Model – Executive Summary"
    body = slide.placeholders[1]
    body.text = (
        "• Identifies early capital rotations across Energy, Defense, AI, and Biotech.\n"
        "• Combines Buffett-style fundamentals with quantitative macro signals.\n"
        "• Proven accuracy: METC (+330%), MP (+320%), BE (+280%), RGTI (+400%).\n"
        "• 30-year compounding horizon anchored on cash-flow quality and moat strength."
    )

    labels = ["AI Infra", "Defense", "Energy", "Materials", "Healthcare", "Catalysts"]
    sizes = [25, 25, 20, 15, 10, 5]
    fig, ax = plt.subplots(figsize=(4, 4))
    ax.pie(
        sizes,
        labels=labels,
        startangle=90,
        colors=["#1f3b73", "#c5a253", "#4169e1", "#708090", "#7fffd4", "#b0c4de"],
    )
    ax.add_artist(plt.Circle((0, 0), 0.55, fc="white"))
    add_chart(prs.slides.add_slide(prs.slide_layouts[6]), fig)

    years = [2025, 2026, 2027, 2028]
    rtx = [7.3, 8.5, 9.5, 10.5]
    lmt = [6.6, 6.9, 7.2, 7.5]
    noc = [3.4, 3.7, 3.8, 4.0]
    fig, ax = plt.subplots()
    ax.plot(years, rtx, "-o", label="RTX")
    ax.plot(years, lmt, "-o", label="LMT")
    ax.plot(years, noc, "-o", label="NOC")
    ax.set_title("Free Cash Flow Outlook (USD B)")
    ax.set_xlabel("Fiscal Year")
    ax.set_ylabel("FCF (B)")
    ax.legend()
    ax.grid(True, ls="--", alpha=0.5)
    add_chart(prs.slides.add_slide(prs.slide_layouts[6]), fig)

    stages = ["Capital", "Atlas Engine", "Fund Returns"]
    values = [100, 80, 60]
    fig, ax = plt.subplots()
    ax.bar(stages, values, color=["#c5a253", "#1f3b73", "#c5a253"])
    ax.set_ylim(0, 110)
    ax.set_title("Capital Deployment Funnel")
    for i, value in enumerate(values):
        ax.text(i, value + 2, f"{value}%", ha="center")
    add_chart(prs.slides.add_slide(prs.slide_layouts[6]), fig)

    return prs


def main() -> None:
    prs = build_presentation()
    prs.save("Atlas_Investor_Deck_v2_Polished.pptx")
    print("Deck created: Atlas_Investor_Deck_v2_Polished.pptx")


if __name__ == "__main__":
    main()
